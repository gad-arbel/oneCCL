"""Build the multi-tab XLSX (one bandwidth chart per metric) for run_occl_pt.

openpyxl is imported lazily inside write_xlsx so it is only required when
--xlsx is used."""

import csv
import sys
from pathlib import Path

# Chart series styling / layout constants (openpyxl uses EMU for line widths).
LINE_WIDTH_EMU = 22225  # 1.75pt series line, in EMU
CHART_WIDTH_CM = 15 * 3  # 3x openpyxl's default chart width
CHART_HEIGHT_CM = 7.5 * 2  # 2x openpyxl's default chart height
# The chart (~15cm tall) spans ~30 rows at the default row height; seat the
# command block this many rows below the chart anchor to clear it.
CMD_BLOCK_ROW_GAP = 34
MERGE_END_COL = 26  # merge the command cell across B..Z so it isn't clipped
CMD_COL = 2  # column B: where each run's command text goes
HYPERLINK_BLUE = "0563C1"


def write_xlsx(xlsx_path, metric_csvs, commands):
    """Build an XLSX from the per-metric pivot CSVs, one tab per metric (e.g.
    'Algorithm Bandwidth', 'Bus Bandwidth', 'Latency'). Each tab holds that
    metric's pivot table at A1 and its line chart below (x = message size, one
    series per run column). Below the chart every run's command is written, one
    per cell, and each run's column header in the table links to its command
    cell. openpyxl is imported on demand so the dependency is only needed when
    --xlsx is used.

    metric_csvs: iterable of (Metric, csv_path).
    commands: list of (column_name, command) accumulated across runs."""
    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("Error: --xlsx requires openpyxl. Install it with "
                 "'pip install openpyxl'.")

    xlsx_path = Path(xlsx_path)
    cmd_by_col = dict(commands)

    workbook = Workbook()
    # Each metric gets its own tab; drop the default empty sheet.
    workbook.remove(workbook.active)
    for metric, csv_path in metric_csvs:
        with Path(csv_path).open(newline="") as f:
            rows = list(csv.reader(f))
        if len(rows) < 2:
            continue  # header only / empty -> no tab, nothing to chart
        sheet = workbook.create_sheet(metric.tab_title)
        _write_table(sheet, rows)
        chart_anchor_row = len(rows) + 2
        _add_chart(sheet, metric, n_rows=len(rows) - 1,
                   run_names=rows[0][1:], anchor_row=chart_anchor_row)
        _write_commands(sheet, rows[0][1:], cmd_by_col,
                        title_row=chart_anchor_row + CMD_BLOCK_ROW_GAP)

    if not workbook.sheetnames:
        # Every metric CSV was header-only, so no tab was created. openpyxl
        # cannot save a sheetless workbook -- skip the write rather than raise.
        print(f"[run_occl_pt] no charted data in any metric; skipping "
              f"{xlsx_path}.", file=sys.stderr)
        return

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)
    print(f"[run_occl_pt] wrote {xlsx_path}")


def _write_table(sheet, rows):
    """Write the pivot CSV at A1. Header text stays string; the size column is
    int; data cells are floats (blank -> None) so the chart plots numbers."""
    for row_i, row in enumerate(rows, start=1):
        for col_i, value in enumerate(row, start=1):
            if row_i == 1:
                cell_value = value
            elif col_i == 1:
                cell_value = int(value)
            else:
                cell_value = float(value) if value else None
            sheet.cell(row=row_i, column=col_i, value=cell_value)


def _add_chart(sheet, metric, n_rows, run_names, anchor_row):
    """Add the metric's line chart below the table: x = message size, one
    series per run column."""
    from openpyxl.chart import LineChart, Reference, Series
    from openpyxl.chart.marker import Marker

    chart = LineChart()
    # Excel auto-enables "vary colors by point" on single-series charts, which
    # colors each segment differently and lists every size in the legend. Force
    # it off so each run stays one color/label.
    chart.varyColors = False
    # Runs need not cover the same sizes (blank cells); connect the line across
    # those gaps instead of breaking it.
    chart.display_blanks = "span"
    chart.title = f"{metric.tab_title} Results"
    chart.x_axis.title = "Data size in bytes"
    chart.y_axis.title = metric.y_label
    chart.style = 10
    # openpyxl defaults leave the axes with no tick marks and no tickLblPos, so
    # Excel/LibreOffice draw the lines but hide the size and bandwidth numbers.
    # Force tick marks + labels, and seat the category axis at the bottom (its
    # default axPos is 'l').
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.x_axis.majorTickMark = "out"
    chart.y_axis.majorTickMark = "out"
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.number_format = "0.00"
    chart.width = CHART_WIDTH_CM
    chart.height = CHART_HEIGHT_CM

    # Build each series explicitly (one per run column), with categories fixed
    # on the size column. The add_data(titles_from_data=True) inference path
    # renders fine in LibreOffice but is read as one-series-per-row by Excel.
    for offset, run_name in enumerate(run_names):
        values = Reference(sheet, min_col=CMD_COL + offset,
                           min_row=2, max_row=1 + n_rows)
        series = Series(values, title=run_name)
        series.marker = Marker(symbol="diamond", size=5)
        series.graphicalProperties.line.width = LINE_WIDTH_EMU
        chart.series.append(series)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=1 + n_rows))
    sheet.add_chart(chart, f"A{anchor_row}")


def _write_commands(sheet, run_names, cmd_by_col, title_row):
    """List each run's command below the chart, one per row, and link each run's
    table header cell to its command cell."""
    from openpyxl.styles import Font

    sheet.cell(row=title_row, column=1, value="Run commands").font = Font(bold=True)
    for offset, run_name in enumerate(run_names):
        cmd_row = title_row + 1 + offset
        sheet.cell(row=cmd_row, column=1, value=run_name).font = Font(bold=True)
        # Commands are long; merge B..Z into one wide cell so the whole command
        # shows without clipping. The value goes in B (top-left of the range).
        sheet.cell(row=cmd_row, column=CMD_COL, value=cmd_by_col.get(run_name, ""))
        sheet.merge_cells(start_row=cmd_row, start_column=CMD_COL,
                          end_row=cmd_row, end_column=MERGE_END_COL)
        # Link the table header cell (row 1, this run's column) to the command
        # cell. Internal refs need the sheet name quoted.
        header_cell = sheet.cell(row=1, column=CMD_COL + offset)
        header_cell.hyperlink = f"#'{sheet.title}'!B{cmd_row}"
        header_cell.font = Font(underline="single", color=HYPERLINK_BLUE)
